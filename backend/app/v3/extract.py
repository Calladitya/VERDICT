"""Safe upload validation + text extraction (PDF, DOCX, TXT, photos/scans)."""
import io
import logging
from typing import Tuple

from fastapi import HTTPException

from . import llm

log = logging.getLogger("verdict.v3.extract")

MAX_BYTES = 10 * 1024 * 1024  # 10 MB
MAX_CHARS = 60_000
ALLOWED = {".pdf", ".docx", ".txt", ".png", ".jpg", ".jpeg", ".webp"}
IMAGE_MIME = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}


def _ext(filename: str) -> str:
    filename = (filename or "").lower()
    return "." + filename.rsplit(".", 1)[-1] if "." in filename else ""


def validate_upload(filename: str, data: bytes) -> str:
    ext = _ext(filename)
    if ext not in ALLOWED:
        raise HTTPException(400, "Unsupported file type. Use PDF, DOCX, TXT, PNG, JPG or WEBP.")
    if not data:
        raise HTTPException(400, "The file is empty.")
    if len(data) > MAX_BYTES:
        raise HTTPException(413, "File is too large (limit 10 MB).")
    head = data[:12]
    ok = (
        (ext == ".pdf" and head.startswith(b"%PDF"))
        or (ext == ".docx" and head.startswith(b"PK"))
        or (ext == ".png" and head.startswith(b"\x89PNG"))
        or (ext in (".jpg", ".jpeg") and head.startswith(b"\xff\xd8"))
        or (ext == ".webp" and head[:4] == b"RIFF" and data[8:12] == b"WEBP")
        or ext == ".txt"
    )
    if not ok:
        raise HTTPException(400, "File content does not match its extension.")
    return ext


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise HTTPException(400, "This PDF is password-protected. Remove the password and upload again.")
    parts = []
    for page in reader.pages[:80]:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            parts.append("")
    return "\n".join(parts).strip()


def _pdf_ocr(data: bytes) -> str:
    """Scanned PDF -> render first pages to images -> vision model. Needs PyMuPDF (optional)."""
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise HTTPException(
            400,
            "This looks like a scanned PDF. Install PyMuPDF (pip install pymupdf) or upload a clear photo (JPG/PNG) instead.",
        )
    doc = fitz.open(stream=data, filetype="pdf")
    texts = []
    for i, page in enumerate(doc):
        if i >= 4:  # keep cost sane: first 4 pages
            break
        pix = page.get_pixmap(dpi=130)
        texts.append(llm.vision_transcribe(pix.tobytes("png"), "image/png"))
    return "\n\n".join(t for t in texts if t).strip()


def _docx_text(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts).strip()


def extract_text(filename: str, data: bytes) -> Tuple[str, str]:
    """Returns (text, method). Raises HTTPException with friendly messages."""
    ext = validate_upload(filename, data)
    method = "text"
    try:
        if ext == ".pdf":
            text = _pdf_text(data)
            method = "pdf"
            if len(text) < 80:
                text = _pdf_ocr(data)
                method = "pdf-ocr"
        elif ext == ".docx":
            text, method = _docx_text(data), "docx"
        elif ext == ".txt":
            text = data.decode("utf-8", errors="replace").strip()
        else:
            text, method = llm.vision_transcribe(data, IMAGE_MIME[ext]), "image-ocr"
    except HTTPException:
        raise
    except Exception as e:  # corrupted file etc.
        log.warning("extract failed for %s: %s", filename, e)
        raise HTTPException(400, "Could not read this file. It may be corrupted.")
    if not text.strip():
        raise HTTPException(400, "No readable text found in this file.")
    return text[:MAX_CHARS], method
