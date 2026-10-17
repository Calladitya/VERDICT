"""DOCX / PDF / ICS generation.

Block format used everywhere:  [("h", "Heading"), ("p", "Paragraph"), ("b", "Bullet")]
"""
import io
import os
import re
from datetime import date, datetime, timedelta
from typing import Iterable, List, Tuple
from xml.sax.saxutils import escape

Block = Tuple[str, str]


def text_to_blocks(text: str) -> List[Block]:
    """Plain text (letter) -> paragraphs. Blank lines separate paragraphs."""
    blocks: List[Block] = []
    for para in re.split(r"\n\s*\n", text.strip()):
        para = para.strip("\n")
        if para:
            blocks.append(("p", para))
    return blocks


# ---------------- DOCX ----------------
def render_docx(title: str, blocks: List[Block], *, show_title: bool = True) -> bytes:
    import docx
    from docx.shared import Pt

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    if show_title and title:
        d.add_heading(title, level=1)
    for kind, text in blocks:
        if kind == "h":
            d.add_heading(text, level=2)
        elif kind == "b":
            d.add_paragraph(text, style="List Bullet")
        else:
            p = d.add_paragraph()
            lines = text.split("\n")
            for i, line in enumerate(lines):
                p.add_run(line)
                if i < len(lines) - 1:
                    p.add_run().add_break()
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---------------- PDF ----------------
_FONT_CACHE = {}


def _pick_font(text: str) -> str:
    """Helvetica for English. For Hindi/Indic text try a Windows/Linux Unicode font."""
    if all(ord(c) < 0x250 for c in text):
        return "Helvetica"
    if "unicode" in _FONT_CACHE:
        return _FONT_CACHE["unicode"]
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        r"C:\Windows\Fonts\Nirmala.ttf",
        r"C:\Windows\Fonts\Nirmala.ttc",
        r"C:\Windows\Fonts\mangal.ttf",
        "/usr/share/fonts/truetype/noto/NotoSansDevanagari-Regular.ttf",
        "/usr/share/fonts/truetype/lohit-devanagari/Lohit-Devanagari.ttf",
    ]
    for path in candidates:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont("V3Unicode", path))
                _FONT_CACHE["unicode"] = "V3Unicode"
                return "V3Unicode"
            except Exception:
                continue
    _FONT_CACHE["unicode"] = "Helvetica"
    return "Helvetica"


def render_pdf(title: str, blocks: List[Block], *, show_title: bool = True) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    all_text = title + " " + " ".join(t for _, t in blocks)
    font = _pick_font(all_text)
    bold = "Helvetica-Bold" if font == "Helvetica" else font
    body = ParagraphStyle("body", fontName=font, fontSize=10.5, leading=15, spaceAfter=6)
    h1 = ParagraphStyle("h1", fontName=bold, fontSize=16, leading=20, spaceAfter=10)
    h2 = ParagraphStyle("h2", fontName=bold, fontSize=12.5, leading=16, spaceBefore=10, spaceAfter=4)
    bullet = ParagraphStyle("bul", parent=body, leftIndent=14, bulletIndent=2)

    story = []
    if show_title and title:
        story.append(Paragraph(escape(title), h1))
    for kind, text in blocks:
        safe = escape(text).replace("\n", "<br/>")
        if kind == "h":
            story.append(Paragraph(safe, h2))
        elif kind == "b":
            story.append(Paragraph(safe, bullet, bulletText="\u2022"))
        else:
            story.append(Paragraph(safe, body))
    story.append(Spacer(1, 6 * mm))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm, bottomMargin=18 * mm,
                      title=title or "VERDICT export").build(story)
    return buf.getvalue()


# ---------------- ICS ----------------
def _ics_escape(s: str) -> str:
    return (s or "").replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\r", "").replace("\n", "\\n")


def _fold(line: str) -> str:
    out, raw = [], line.encode("utf-8")
    while len(raw) > 74:
        cut = 74
        while cut > 0 and (raw[cut] & 0xC0) == 0x80:  # don't split UTF-8 chars
            cut -= 1
        out.append(raw[:cut].decode("utf-8"))
        raw = b" " + raw[cut:]
    out.append(raw.decode("utf-8"))
    return "\r\n".join(out)


def make_ics(events: Iterable[dict], cal_name: str = "VERDICT deadlines") -> str:
    """events: dicts with id, title, detail, date (date obj). All-day events with a 1-day-before reminder."""
    lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//VERDICT//Deadlines//EN",
        "CALSCALE:GREGORIAN",
        f"X-WR-CALNAME:{_ics_escape(cal_name)}",
    ]
    stamp = datetime.utcnow().strftime("%Y%m%dT%H%M%SZ")
    for e in events:
        d: date = e["date"]
        lines += [
            "BEGIN:VEVENT",
            f"UID:{e['id']}@verdict.local",
            f"DTSTAMP:{stamp}",
            f"DTSTART;VALUE=DATE:{d.strftime('%Y%m%d')}",
            f"DTEND;VALUE=DATE:{(d + timedelta(days=1)).strftime('%Y%m%d')}",
            f"SUMMARY:{_ics_escape(e['title'])}",
            f"DESCRIPTION:{_ics_escape(e.get('detail', ''))}",
            "BEGIN:VALARM",
            "ACTION:DISPLAY",
            "DESCRIPTION:Deadline tomorrow",
            "TRIGGER:-P1D",
            "END:VALARM",
            "END:VEVENT",
        ]
    lines.append("END:VCALENDAR")
    return "\r\n".join(_fold(l) for l in lines) + "\r\n"
