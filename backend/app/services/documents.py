"""Document text extraction for PDF / DOCX / TXT uploads."""
from __future__ import annotations

import io


class ExtractionError(Exception):
    pass


def extract_text(filename: str, content: bytes, content_type: str = "") -> str:
    name = filename.lower()
    try:
        if name.endswith(".pdf") or content_type == "application/pdf":
            return _extract_pdf(content)
        if name.endswith(".docx") or "word" in content_type:
            return _extract_docx(content)
        if name.endswith((".txt", ".md")) or content_type.startswith("text/"):
            return content.decode("utf-8", errors="replace")
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"Failed to read document: {exc}") from exc
    raise ExtractionError(
        "Unsupported file type. Please upload a PDF, DOCX, or TXT file."
    )


def _extract_pdf(content: bytes) -> str:
    from pypdf import PdfReader
    print("PDF BYTES RECEIVED:", len(content))
    reader = PdfReader(io.BytesIO(content))
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n\n".join(pages).strip()
    if not text:
        raise ExtractionError(
            "No selectable text found. This may be a scanned PDF that needs OCR "
            "(OCR is planned for a later version)."
        )
    return text


def _extract_docx(content: bytes) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    text = "\n".join(parts).strip()
    if not text:
        raise ExtractionError("The document appears to be empty.")
    return text
